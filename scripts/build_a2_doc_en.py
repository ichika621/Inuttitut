# -*- coding: utf-8 -*-
"""English edition of the A2 example-sentence document, written for review by an Inuttitut
speaker or teacher. It is generated from exactly the data the game uses (a2_bank.py,
a2_bank_en.py, rpg_content.py), so it contains every Inuttitut string the game can show:
  - all examples and word-list entries (with the IDs used in the game),
  - every conversation, with each reply marked as accepted or not,
  - every grammar drill with its answer options,
  - every other Inuttitut word or word-part (glossary).
Usage: python scripts/build_a2_doc_en.py [out.docx]   (needs: pip install python-docx)
"""
import sys, os, re, collections
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import a2_bank as B
import a2_bank_en as E
import rpg_content as C
from build_rpg_lexicon import lexicon

LEXD = lexicon()
ITEMS = LEXD["items"]
BY_INU = {it["inu"]: it for it in ITEMS}
errs = C.validate(set(BY_INU))
if errs:
    raise SystemExit("Game content check failed:\n  " + "\n  ".join(errs))

NAVY = "1F3A5F"; AMBER = (0x9A, 0x5B, 0x13); GREY = (0x55, 0x55, 0x55)
FILL = {"P": "E8F0E3", "T": "E3ECF5", "R": "F6EEDC", "W": "F3E3E3"}
SRC_SHORT = {"I1": "Course I #1", "I2": "Course I #2", "I3": "Course I #3", "I8": "Course I #8", "I9": "Course I #9", "I10": "Course I Numbers",
             "I11": "Course I #11", "II1": "Course II #1", "II3": "Course II #3", "II8": "Course II #8", "PT": "Practice Test",
             "NG": "Nunatsiavut Gov.", "TU2": "Tusâlanga L2", "TU3": "Tusâlanga L3", "TU4": "Tusâlanga L4", "TU5": "Tusâlanga L5",
             "TUT": "Tusâlanga Gr.41", "TUP": "Tusâlanga Gr.29", "DIC": "LVM dictionary"}
SPEAKER = {k: re.sub(r"\{\{(.+?)\}\}", r"\1", v["en"]) for k, v in C.NPC_NAMES.items()}
SCRIPT_TITLE = {
    "intro": "Opening (narration only)", "greet_anansiak": "Quest 1 — Greet Grandma", "greet_hulda": "Quest 1 — Greet Hulda",
    "greet_tami": "Quest 1 — Greet Tâmi", "hulda_intro": "Quest 2 — Hulda: where are you from?", "baby": "Quest 2 — Grandma: the baby (before lesson 4)",
    "tami_coffee": "Quest 3 — Tâmi: coffee", "tami_dogs": "Quest 3 — Tâmi: how many dogs?", "hulda_stars": "Quest 4 — Hulda: counting stars",
    "rita_talk": "Quest 5 — Rita at the shore", "makusi_weather": "Quest 6 — Makusi: the weather", "makusi_time": "Quest 6 — Makusi: the plane's time",
    "ending": "Ending — Grandma"}

# where each example is used in conversations
USED = collections.defaultdict(set)
for key, steps in list(C.SCRIPTS.items()) + [("small talk", s) for s in C.SMALLTALK.values()]:
    for st in steps:
        if "say" in st: USED[st["say"]].add("dialogue")
        for c in st.get("choose", []): USED[c["t"]].add("reply")
for o in C.HOME_OPTIONS: USED[o["t"]].add("reply")
DRILL_FULL = {}
for u, ds in C.DRILLS.items():
    for d in ds: DRILL_FULL["".join(d.split("|"))] = u
for sec in C.ABOUT:
    for b in sec.get("breakdowns", []): USED[b["word"]].add("about page")
    for _, ws in sec.get("examples", []):
        for w in ws: USED[w].add("about page")

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(1.8)
sec.left_margin = sec.right_margin = Cm(1.9)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)


def shade(cell, hex_):
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
    cell._tc.get_or_add_tcPr().append(shd)


def rich(par, text, size=None, color=None, bold=False, italic=False):
    """Write text, rendering {{word}} marks as bold italics."""
    for i, chunk in enumerate(re.split(r"\{\{(.+?)\}\}", text)):
        if not chunk: continue
        r = par.add_run(chunk)
        r.bold = bold or i % 2 == 1; r.italic = italic or i % 2 == 1
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor(*color)


def p(text="", bold=False, italic=False, size=None, color=None, after=4):
    para = doc.add_paragraph(); rich(para, text, size, color, bold, italic); para.paragraph_format.space_after = Pt(after); return para


def bullet(text, size=10):
    para = doc.add_paragraph(style="List Bullet"); rich(para, text, size); para.paragraph_format.space_after = Pt(1)


def table(headers, rows, widths, size=9, fill=NAVY, fills=None, bold_cols=()):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.autofit = False
    for i, (h, w) in enumerate(zip(headers, widths)):
        c = t.rows[0].cells[i]; c.width = Cm(w); c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, fill); c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, (v, w) in enumerate(zip(row, widths)):
            cells[i].width = Cm(w); cells[i].text = ""
            rich(cells[i].paragraphs[0], str(v), size, bold=i in bold_cols)
            if fills and fills[ri]: shade(cells[i], fills[ri])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def toc():
    para = doc.add_paragraph(); r = para.add_run()._r
    for kind, val in (("begin", None), ("instr", 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), ("text", "(Table of contents: in Word, right-click → Update Field)"), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = val
        elif kind == "text":
            el = OxmlElement("w:t"); el.text = val
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        r.append(el)


def src_cell(it):
    s = E.KIND_SHORT_EN[it["kind"]] + " · " + SRC_SHORT.get(it["src"], it["src"])
    return s


cnt = collections.Counter(it["kind"] for it in ITEMS)

# ------------------------------------------------------------ cover
doc.add_heading("Labrador Inuttitut — A2 Example Sentences", level=0)
p("Nunatsiavummiutut / Inuttut · example sentences from A1 to A2, with every Inuttitut line used in the game 'Words of the Snow Village'", italic=True, size=11, color=GREY, after=10)
table(["", ""], [
    ["Language", "Inuttitut (Inuttut, Nunatsiavummiutut), the Inuit language of Nunatsiavut, northern Labrador, Canada. Labrador standard spelling (1980): â = aa, e = ii, o = uu, K = uvular."],
    ["Goal", "CEFR A2: handling simple, everyday exchanges on familiar topics."],
    ["Contents", f"12 units, {len([i for i in ITEMS if '-v' not in i['id']])} example sentences and {len([i for i in ITEMS if '-v' in i['id']])} family words — {len(ITEMS)} entries in all, each with the ID used in the game. Appendices: all game conversations, grammar drills, and other Inuttitut words."],
    ["Where the Inuttitut comes from", f"Printed in the course handouts {cnt['P']} · course paradigm tables {cnt['T']} · built only from course forms and rules {cnt['R']} · web sources {cnt['W']}"],
    ["Main sources", "INUTTITUT I & II course handouts (Katimanik #1–#11, Sept.–Nov. 2025) and Practice Test (13 Nov. 2025); Nunatsiavut Government language page; Tusâlanga (Nunatsiavummiutut); Labrador Virtual Museum English–Inuttut Dictionary."],
], [4.0, 13.4], size=9.5, fill="FFFFFF")
doc.add_heading("Contents", level=1)
toc()
doc.add_page_break()

# ------------------------------------------------------------ 0. reviewer note
doc.add_heading("0. For the reviewer", level=1)
p("Thank you for checking this. This document contains every piece of Inuttitut that appears in the game. The game is built from the same data, and it refuses to start if a line is missing from this document. So if everything here is correct, the Inuttitut in the game is correct.", after=6)
p("What to check", bold=True, color=AMBER, after=2)
for b in [
    "Sections 4 and 5: is each example spelled correctly, and does the English say what it means? Each entry has an ID (for example u3-02). Please use the ID when you note a correction.",
    "Entries marked 'Course rules' were not printed in the handouts; they were put together from forms and rules that were. These especially need a check.",
    "Appendix A: in each conversation, is the line natural for that situation? Are the replies marked ✓ (accepted) really good answers, and the ones marked ✗ really wrong?",
    "Appendix B: each grammar drill shows a word with a gap; the learner picks the missing part. Check that only the marked answer is correct.",
    "Appendix C: other Inuttitut words and word parts used in narration and on the 'What is Inuttitut?' page.",
    "The game does not use any synthesised voice. Pronunciation can only be added as real recordings by a speaker.",
]:
    bullet(b)

# ------------------------------------------------------------ 1. about
doc.add_heading("1. What is Inuttitut?", level=1)
p("This is the same text the game shows on its start screen. Letters in brackets refer to the sources at the end of this section.", italic=True, size=9.5, color=GREY)
for s in C.ABOUT:
    doc.add_heading(s["h"]["en"], level=2)
    for x in s["p"]: p(x["en"])
    for b in s.get("breakdowns", []):
        it = BY_INU[b["word"]]
        para = doc.add_paragraph(); r = para.add_run(b["word"]); r.bold = True; r.font.size = Pt(12)
        para.add_run("   = " + it["en"]).italic = True
        table(["Part", "Meaning", "Source"], [[pt, C.GLOSSARY[pt][1]["en"], C.GLOSSARY[pt][2]] for pt in b["parts"]], [3.2, 8.2, 6.0], size=9)
    for intro, words in s.get("examples", []):
        p(intro["en"], after=1)
        table(["ID", "Inuttitut", "English"], [[BY_INU[w]["id"], w, BY_INU[w]["en"]] for w in words], [2.0, 6.0, 9.4], size=9, bold_cols=(1,))
    for x in s.get("p2", []): p(x["en"])
p("Sources", bold=True, color=AMBER, after=1)
for sid, text, url in C.ABOUT_SOURCES:
    bullet(f"[{sid}] {text}" + (f" — {url}" if url else ""), 9.5)

# ------------------------------------------------------------ 2. sources
doc.add_heading("2. Where the examples come from", level=1)
p("The course handouts are treated as the most reliable source, followed by Labrador-dialect web sources. Standard Inuktitut (Nunavut and elsewhere) is not used, because its endings differ (for example, 'I see' is takujunga in standard Inuktitut but takuvunga in the course).", after=6)
table(["Label", "Meaning"], [
    ["Course: text", "Printed in the course handouts, copied exactly"],
    ["Course: table", "A form printed in one of the course's paradigm tables"],
    ["Course rules", "Built only from forms and rules printed in the course (for example, the answer to a Practice Test question). The basis is given in a note."],
    ["Web", "From a Labrador-dialect web source; the URL is listed below. Only used when it agrees with the course."],
], [3.2, 14.2], size=9, fills=[FILL["P"], FILL["T"], FILL["R"], FILL["W"]])
table(["Code", "Source"], [[SRC_SHORT.get(k, k), v] for k, v in B.SRC.items()], [3.2, 14.2], size=8.5)
p("Handwritten notes on the handouts (the learner's own translations) are not used as a source. The English translations were written from the course's grammar explanations, the dictionary, and the meanings of the suffixes.", italic=True, size=9, color=GREY)

# ------------------------------------------------------------ 3. spelling
doc.add_heading("3. Spelling and pronunciation (course handout I #1)", level=1)
table(["Item", "Spelling", "Examples"], [
    ["Short vowels", "a / i / u", "a as in 'father'; i like 'ee' in 'feet' but shorter; u like 'oo' in 'boot' but shorter"],
    ["Long vowels", "â (= aa) / e (= ii) / o (= uu)", "annâk (two women), angutek (two men), inok (two people)"],
    ["Vowel pairs", "ai, au, ia, iu, ua, ui", "taima / auka 'no' / niaKuk 'head' / siutik 'ear' / pualuk 'mitten' / puijik 'seal'"],
    ["k and K", "k is like English k; K is made further back in the throat (uvular)", "kimmik 'heel' and Kimmik 'dog' are different words"],
    ["ng, long ng", "underlined ng = a long ng", "angutik 'man'; pinnguak / pinguak 'toy' (spelling varies)"],
    ["Consonant pairs", "tl / dl, ts, tj / dj", "atlak (adlak) 'black bear', Nunatsiavut, utjuk (udjuk) 'square flipper seal'"],
], [3.0, 5.4, 9.0], size=9)

# ------------------------------------------------------------ 4. path
doc.add_heading("4. The path to A2", level=1)
table(["#", "Unit", "CEFR", "You will be able to…"],
      [[i + 1, E.UNIT_EN[u["id"]]["title"], u["cefr"], E.UNIT_EN[u["id"]]["cando"]] for i, u in enumerate(B.UNITS)], [0.8, 6.0, 1.4, 9.2], size=9)
doc.add_page_break()

# ------------------------------------------------------------ 5. units
doc.add_heading("5. Examples by unit", level=1)
p("Column 'In game': L = lesson card and word book (every entry); D = said by a character in a conversation; R = offered as a reply choice; G = grammar drill; A = 'What is Inuttitut?' page.", italic=True, size=9, color=GREY)
for ui, u in enumerate(B.UNITS):
    en = E.UNIT_EN[u["id"]]
    doc.add_heading(f"Unit {ui + 1}  {en['title']}  ({u['cefr']})", level=2)
    para = doc.add_paragraph(); r = para.add_run("You will be able to: "); r.bold = True; para.add_run(en["cando"])
    p("Grammar points", bold=True, size=10.5, color=AMBER, after=1)
    for h, body in en["grammar"]:
        para = doc.add_paragraph(style="List Bullet"); rr = para.add_run(h + ": "); rr.bold = True; rr.font.size = Pt(10)
        r2 = para.add_run(body); r2.font.size = Pt(10); para.paragraph_format.space_after = Pt(1)
    p("Examples", bold=True, size=10.5, color=AMBER, after=1)
    rows, fills, notes = [], [], []
    for it in [x for x in ITEMS if x["unit"] == u["id"]]:
        mark = ""
        if it["note"]:
            notes.append((it["id"], it["note"]["en"])); mark = " *"
        used = ["L"] + (["D"] if "dialogue" in USED[it["inu"]] else []) + (["R"] if "reply" in USED[it["inu"]] else []) + \
               (["G"] if it["inu"] in DRILL_FULL else []) + (["A"] if "about page" in USED[it["inu"]] else [])
        rows.append([it["id"], it["inu"], it["en"], src_cell(it) + mark, " ".join(used)])
        fills.append(FILL[it["kind"]])
    table(["ID", "Inuttitut", "English", "Source", "In game"], rows, [1.7, 5.0, 5.3, 3.9, 1.5], size=8.8, fills=fills, bold_cols=(1,))
    for iid, n in notes:
        p(f"* {iid}: {n}", size=8.5, color=GREY, after=1)

# ------------------------------------------------------------ 6. grammar tables
doc.add_page_break()
doc.add_heading("6. Grammar tables (from the course handouts)", level=1)
doc.add_heading("6.1 Intransitive person endings (I #8, I #11)", level=2)
table(["", "1st person", "2nd person", "3rd person"], [["Singular (1)", "-vunga", "-vutit", "-juk"], ["Dual (2)", "-vuguk", "-vutik", "-jok"], ["Plural (3+)", "-vugut", "-vusi", "-jut"]], [3.0, 4.8, 4.8, 4.8])
doc.add_heading("6.2 Negative endings (I #8, I #11)", level=2)
table(["", "1st person", "2nd person", "3rd person"], [["Singular", "-ngilanga", "-ngilatit", "-ngituk"], ["Dual", "-ngilaguk", "-ngilatik", "-ngitok"], ["Plural", "-ngilagut", "-ngilasi", "-ngitut"]], [3.0, 4.8, 4.8, 4.8])
p("In conversation the 3rd-person singular also appears as -ngilak (I #9: Anugingilak.).", size=9, color=GREY)
doc.add_heading("6.3 Question endings (I #8, I #11)", level=2)
table(["", "1st person", "2nd person", "3rd person"], [["Singular", "-vingâ?", "-ven?", "-jong?"], ["Dual", "-vinong?", "-viteng?", "-joung?"], ["Plural", "-vitâ?", "-vise?", "-jon?"]], [3.0, 4.8, 4.8, 4.8])
doc.add_heading("6.4 Commands (I #8; written -get / -liget / -lauget in I #11)", level=2)
table(["To", "Endings", "Examples"], [["one person", "-git / -ligit / -laugit", "Tupagit! / Tupaligit! / Tupalaugit!"], ["two people", "-gittek / -ligittek / -laugittek", "Tupagittek! / Tupaligittek! / Tupalaugittek!"], ["three or more", "-gitsi / -ligitsi / -laugitsi", "Tupagitsi! / Tupaligitsi! / Tupalaugitsi!"]], [2.8, 5.8, 8.8])
doc.add_heading("6.5 Possessive endings (I #8, I #11)", level=2)
table(["Owned", "my", "your", "his / her"], [["one thing", "-ga", "-it", "-nga"], ["two things", "-Vkka", "-Vkkik", "-Vgik"], ["three or more", "-kka", "-tit", "-ngit"]], [3.0, 4.8, 4.8, 4.8])
table(["Noun", "my (1 / 2 / 3+)", "your (1 / 2 / 3+)", "her (1 / 2 / 3+)"], [
    ["panik daughter", "paniga / panekka / panikka", "panet / panekkik / panitit", "paninga / panegik / paningit"],
    ["innik son", "inniga / inneka / innika", "innet / innekik / innitit", "inninga / innegik / inningit"],
    ["illuk house", "illuga / illoka / illuka", "illuit / illokik / illutit", "illunga / illogik / illungit"],
    ["Kajak kayak", "Kajaga / Kajâkka / Kajakka", "Kajait / Kajâkkik / Kajatit", "Kajanga / Kajâgik / Kajangit"],
    ["umiak boat", "umiaga / umiakka / umiakka", "umiat / umiakkik / umiatit", "umianga / umiagik / umiangit"]], [3.0, 4.8, 4.8, 4.8], size=8.8)
p("These full possessive paradigms are printed in course handout I #8. They appear in the game only as listed in Unit 6.", size=9, color=GREY)
doc.add_heading("6.6 Noun cases (II #1, I #11)", level=2)
table(["Role", "Singular", "Dual", "Plural", "Meaning"], [
    ["subject (pijuk)", "-Ø", "-V", "-it", "—"], ["object (pijak)", "-mik", "-Vnik", "-nik", "(object)"], ["owner", "-up", "-V", "-it", "'s / of"],
    ["direction", "-mut", "-Vnut", "-nut", "to"], ["origin", "-mit", "-Vnit", "-nit", "from"], ["place", "-mi", "-Vni", "-ni", "in / at"],
    ["route / means", "-kkut", "-Vkkut", "-tigut", "through / by"], ["manner", "-tut", "-Vtut", "-titut", "like"]], [3.4, 2.6, 2.6, 2.6, 6.2], size=9)
p("V = double the final vowel of the root (course handout note).", size=9, color=GREY)
doc.add_heading("6.7 Tense (I #9, pisuk- 'walk')", level=2)
table(["Tense", "Suffix", "I", "you", "he / she", "I (negative)"], [
    ["before today", "-lauk-", "pisulaukKunga", "pisulaukKutit", "pisulauttuk", "pisulaungilanga"],
    ["earlier today", "-kKau-", "pisukKauvunga", "pisukKauvutit", "pisukKaujuk", "pisukKaungilanga"],
    ["just now", "-tainnak-", "pisutainnaKunga", "pisutainnaKutit", "pisutainnatuk", "pisutainnangilanga"],
    ["present", "—", "pisuvunga", "pisuvutit", "pisujuk", "pisungilanga"],
    ["right now", "-lik- / -lu(k)-", "pisulikKunga", "pisulikKutit", "pisulittuk", "pisulungilanga"],
    ["about to", "-giasi-", "pisugiasivunga", "pisugiasivutit", "pisugiasijuk", "pisugiasingngilanga"],
    ["later today", "-niak-", "pisuniakKunga", "pisuniakKutit", "pisuniattuk", "pisuniangilanga"],
    ["later today (Hebron)", "-langa- / -langalu-", "pisulangavunga", "pisulangavutit", "pisulangajuk", "pisulangalungilanga"],
    ["tomorrow or later", "-lâk-", "pisulâkKunga", "pisulâkKutit", "pisulâttuk", "pisulângilanga"]], [3.0, 2.8, 3.0, 2.9, 2.8, 3.1], size=8.3)
p("This table is printed in handout I #9; the game uses the forms listed in Units 3, 5 and 8.", size=9, color=GREY)
doc.add_heading("6.8 Transitive takuk- 'see' (II #3)", level=2)
table(["object \\ subject", "I", "you", "he / she"], [["me", "—", "takuvamma", "takuvânga"], ["you", "takuvagit", "—", "takuvâtit"], ["him / her", "takujaga", "takujait", "takujanga"]], [3.6, 4.6, 4.6, 4.6])
table(["object \\ subject (question)", "I", "you", "he / she"], [["me", "—", "takuvammâ?", "takuvângâ?"], ["you", "takuvagen?", "—", "takuvâten?"], ["him / her", "takujagâ?", "takujân?", "takujaung?"]], [3.6, 4.6, 4.6, 4.6])
doc.add_heading("6.9 Derivational suffixes (I #8, I #11)", level=2)
table(["Type", "Suffixes"], [
    ["noun → verb", "-u- (be), -Ka- (have), -lia- (go to), -liu- (make), -tu- (eat / drink), -siu- (look for)"],
    ["verb → verb", "-titsi- (make someone do), -guma- (want to), -songu- (can), -giaKa- (have to), -Kattak- (habitually), -gunna- (may), -giattu- (go to do), -gasua- (try to), -jutsau- (should be), -galâ- (a little), -kasâ- (almost), -gajak- (could), -galuak- (although)"],
    ["noun → noun", "-nguak (toy / pretend), -tsiak (good / beautiful), -miuk (person from), -vinik (meat of)"],
    ["verb → noun", "-vik (place for)"]], [3.0, 14.4], size=9)
doc.add_heading("6.10 Numbers", level=2)
table(["", "1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "11", "12"], [
    ["Numbers (I)", "atautsik", "maggok", "pingasut", "sitamat", "tallimat", "pingasojuttut", "—", "sitamaujuttut", "—", "tallimaujuttut", "—", "—"],
    ["Clock (II #8)", "ainsik", "suvait", "tarait", "fiarat", "fimfit", "sâtsit", "sepat", "âttat", "nainat", "senat", "ailfat", "suvailfat"]], [2.2] + [1.27] * 12, size=7.5)
p("— = not given in the handout (7 and 9 are blank in the course numbers list; 11 and 12 are not listed there).", size=9, color=GREY)

# ------------------------------------------------------------ Appendix A: dialogues
doc.add_page_break()
doc.add_heading("Appendix A. Every conversation in the game", level=1)
p("Each conversation is listed in the order the player sees it. Narration is in English in the English version of the game (and Japanese in the Japanese version); Inuttitut words inside narration are shown in bold italics and are listed in Appendix C. ✓ = the game accepts this reply; ✗ = the game says it's wrong and shows the hint.", italic=True, size=9.5, color=GREY)


def dump_steps(steps, default_who=None):
    last = default_who
    for stp in steps:
        if "say" in stp:
            last = stp["who"]; it = BY_INU[stp["say"]]
            para = doc.add_paragraph(); r = para.add_run(SPEAKER[stp["who"]] + ": "); r.bold = True
            r2 = para.add_run(stp["say"]); r2.bold = True; r2.font.size = Pt(11.5)
            para.add_run(f"   ({it['id']}) “{it['en']}”").italic = True
            para.paragraph_format.space_after = Pt(1)
        elif "narr" in stp:
            p("[narration] " + stp["narr"]["en"], size=9.5, color=GREY, after=1)
        elif "pic" in stp:
            p({"stars4": "[picture: four stars in the night sky]", "sunny": "[picture: a sunny, calm day]", "clock1030": "[picture: a clock at 10:30]"}[stp["pic"]], size=9.5, color=GREY, after=1)
        elif "choose" in stp or "chooseHome" in stp:
            opts = stp.get("choose")
            if stp.get("chooseHome"):
                opts = [{"t": o["t"], "ok": 1 if o["home"] else 0, "hint": o["hint"], "_home": o["home"]} for o in C.HOME_OPTIONS]
            rows = []
            for o in opts:
                it = BY_INU[o["t"]]
                acc = "✓" if o.get("ok") else "✗"
                if o.get("_home"): acc = f"✓ if the player said earlier they are from {o['_home']}"
                rows.append([acc, o["t"], it["en"], it["id"], (o.get("hint") or {}).get("en", "")])
            p("Player's reply choices:", size=9.5, bold=True, after=1)
            table(["", "Inuttitut", "English", "ID", "Hint shown if chosen"], rows, [1.6, 4.6, 4.2, 1.6, 5.4], size=8.6, bold_cols=(1,))
        elif "quiz" in stp:
            p("[comprehension question, answered in English/Japanese] " + stp["quiz"]["en"] + "  —  " + " / ".join(("✓ " if i == stp["ans"] else "") + o["en"] for i, o in enumerate(stp["opts"])), size=9.5, color=GREY, after=1)


for key, steps in C.SCRIPTS.items():
    doc.add_heading(SCRIPT_TITLE.get(key, key), level=2)
    dump_steps(steps)
doc.add_heading("Small talk (when there is no quest conversation)", level=2)
for npc, steps in C.SMALLTALK.items():
    p(SPEAKER[npc], bold=True, color=AMBER, after=1)
    dump_steps(steps, npc)

# ------------------------------------------------------------ Appendix B: drills
doc.add_page_break()
doc.add_heading("Appendix B. Grammar drills", level=1)
p("The learner sees the English meaning and the sentence with a gap (＿＿), and picks the missing part. The options offered are the other parts from the same drill set (up to four in total). The full sentence is always one of the examples in Section 5.", italic=True, size=9.5, color=GREY)
for u, ds in C.DRILLS.items():
    targets = []
    for d in ds:
        t_ = d.split("|")[1]
        if t_ not in targets: targets.append(t_)
    doc.add_heading(f"{E.UNIT_EN[u]['title']}", level=2)
    p("Parts that can appear as options: " + ", ".join("-" + x + "-" for x in targets), size=9.5, after=2)
    rows = []
    for d in ds:
        stem, target, rest = d.split("|"); full = stem + target + rest; it = BY_INU[full]
        rows.append([it["id"], stem + "＿＿" + rest, "-" + target + "-", full, it["en"]])
    table(["ID", "Shown", "Answer", "Full sentence", "English"], rows, [1.6, 4.2, 2.4, 4.4, 4.8], size=8.8, bold_cols=(3,))

# ------------------------------------------------------------ Appendix C: glossary
doc.add_heading("Appendix C. Other Inuttitut in the game", level=1)
p("Words and word parts that appear outside the examples: in narration, quest text, character names, and the 'What is Inuttitut?' page. Personal names are from the course dialogues. Place names (Nain, Hopedale, Makkovik, Postville, Rigolet, Hebron) are written in their English form in the game's narration.", italic=True, size=9.5, color=GREY)
table(["Word / part", "Kind", "Meaning", "Source"],
      [[k, {"word": "word", "part": "word part", "name": "name"}[v[0]], v[1]["en"], v[2]] for k, v in C.GLOSSARY.items()], [3.4, 2.0, 7.0, 5.0], size=8.8, bold_cols=(0,))
p("The game's title screen also shows Ilitannamek! (u1-01).", size=9.5, color=GREY)

# ------------------------------------------------------------ Appendix D: coverage
doc.add_heading("Appendix D. Coverage check", level=1)
n_dialog = len({k for k, v in USED.items() if "dialogue" in v})
n_reply = len({k for k, v in USED.items() if "reply" in v})
table(["What", "Count"], [
    ["Entries in the example bank (all shown in lessons and the word book)", len(ITEMS)],
    ["…of which spoken by characters in conversations", n_dialog],
    ["…of which offered as reply choices", n_reply],
    ["Grammar drills", sum(len(v) for v in C.DRILLS.values())],
    ["Conversations (including small talk)", len(C.SCRIPTS) + len(C.SMALLTALK)],
    ["Other Inuttitut words and parts (Appendix C)", len(C.GLOSSARY)],
], [13.0, 4.4], size=9.5)
p("The build script checks automatically that every Inuttitut line in the game's conversations, replies, drills, narration marks and the 'What is Inuttitut?' page appears in this document; if not, the game data is not built.", size=9.5)

OUT = sys.argv[1] if len(sys.argv) > 1 else "Inuttitut_A2_Example_Sentences_EN.docx"
doc.save(OUT)
print("saved", OUT, len(ITEMS), "entries")
